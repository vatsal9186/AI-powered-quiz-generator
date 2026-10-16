import io
import re
import pypdf
import docx
import pptx
import olefile

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extracts text from PDF bytes using pymupdf, inserting page markers.
    """
    import pymupdf
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    text_parts = []
    for i, page in enumerate(doc):
        page_text = page.get_text() or ""
        text_parts.append(f"\n--- Page {i+1} ---\n{page_text}")
    doc.close()
    return "\n".join(text_parts)

def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extracts text from DOCX bytes, including paragraphs and tables.
    """
    docx_file = io.BytesIO(file_bytes)
    doc = docx.Document(docx_file)
    text_parts = []
    
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
            
    for table in doc.tables:
        for row in table.rows:
            # Gather unique cell texts in order to avoid repeating merged cells multiple times
            row_text = []
            for cell in row.cells:
                ct = cell.text.strip()
                if ct and (not row_text or row_text[-1] != ct):
                    row_text.append(ct)
            if row_text:
                text_parts.append(" | ".join(row_text))
                
    return "\n".join(text_parts)

def extract_text_from_pptx(file_bytes: bytes) -> str:
    """
    Extracts text from PPTX bytes, slide-by-slide.
    """
    pptx_file = io.BytesIO(file_bytes)
    prs = pptx.Presentation(pptx_file)
    text_parts = []
    
    for i, slide in enumerate(prs.slides):
        text_parts.append(f"\n--- Slide {i+1} ---")
        if slide.shapes.title:
            text_parts.append(f"Title: {slide.shapes.title.text}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    para_text = "".join(run.text for run in paragraph.runs).strip()
                    if para_text:
                        text_parts.append(para_text)
                        
    return "\n".join(text_parts)

def extract_text_from_doc(file_bytes: bytes) -> str:
    """
    Extracts readable text sequences from legacy DOC binary OLE structures.
    """
    try:
        with olefile.OleFileIO(io.BytesIO(file_bytes)) as ole:
            if ole.exists("WordDocument"):
                stream = ole.openstream("WordDocument")
                data = stream.read()
                # Find ASCII printable sequences of length >= 4
                words = re.findall(r'[\x20-\x7E\x0A\x0D]{4,}', data.decode('ascii', errors='ignore'))
                return "\n".join(words)
    except Exception:
        pass
    # Standard printable characters fallback
    words = re.findall(r'[\x20-\x7E\x0A\x0D]{4,}', file_bytes.decode('ascii', errors='ignore'))
    return "\n".join(words)

def extract_text_from_ppt(file_bytes: bytes) -> str:
    """
    Extracts readable text sequences from legacy PPT binary OLE structures.
    """
    try:
        with olefile.OleFileIO(io.BytesIO(file_bytes)) as ole:
            if ole.exists("PowerPoint Document"):
                stream = ole.openstream("PowerPoint Document")
                data = stream.read()
                # Find ASCII printable sequences of length >= 4
                words = re.findall(r'[\x20-\x7E\x0A\x0D]{4,}', data.decode('ascii', errors='ignore'))
                return "\n".join(words)
    except Exception:
        pass
    # Standard printable characters fallback
    words = re.findall(r'[\x20-\x7E\x0A\x0D]{4,}', file_bytes.decode('ascii', errors='ignore'))
    return "\n".join(words)

def extract_text_for_validated_file(file_bytes: bytes, detected_format: str) -> str:
    """
    Dispatches to the correct parser based on the validated format.
    """
    if detected_format == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif detected_format == "docx":
        return extract_text_from_docx(file_bytes)
    elif detected_format == "pptx":
        return extract_text_from_pptx(file_bytes)
    elif detected_format == "doc":
        return extract_text_from_doc(file_bytes)
    elif detected_format == "ppt":
        return extract_text_from_ppt(file_bytes)
    else:
        raise ValueError(f"Unknown validated format: {detected_format}")
